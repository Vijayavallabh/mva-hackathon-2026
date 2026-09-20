#!/usr/bin/env python3
"""Pinned public Evo2 weights and public-reference-only DNA benchmark inputs."""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import gzip
import hashlib
import json
from pathlib import Path
import urllib.request
from track2_esm_pilot import digest, write_json

ROOT=Path('/home/prachh/v/mva-track2-expanded-20260920')


def weights(name):
    from huggingface_hub import HfApi, hf_hub_download
    if name not in ['evo2_7b','evo2_40b']:raise ValueError('Unregistered checkpoint')
    repo='arcinstitute/'+name
    info=HfApi(token=False).model_info(repo,files_metadata=True)
    files=[s for s in info.siblings if s.rfilename==name+'.pt' or s.rfilename.startswith(name+'.pt.part')]
    if not files or sum(s.size for s in files)>100_000_000_000:
        raise ValueError('Unexpected checkpoint files/size')
    dest=ROOT/'cache/evo2'/name;dest.mkdir(parents=True,exist_ok=False)
    manifest={'repository':repo,'revision':info.sha,'created_utc':datetime.now(timezone.utc).isoformat(),'files':[]}
    for item in files:
        path=Path(hf_hub_download(repo,item.rfilename,revision=info.sha,local_dir=dest,token=False))
        sha=digest(path)
        upstream=item.lfs.sha256 if item.lfs else None
        if path.stat().st_size!=item.size or (upstream and sha!=upstream):
            raise ValueError('Weight differs from pinned upstream size/LFS digest')
        manifest['files'].append(dict(name=path.name,bytes=item.size,sha256=sha,upstream_lfs_sha256=upstream))
    target=dest/(name+'.pt')
    if not target.exists():
        parts=[dest/f'{name}.pt.part{i}' for i in range(len(files))]
        if not all(p.exists() for p in parts):raise ValueError('Noncontiguous checkpoint parts')
        with target.open('xb') as out:
            for p in parts:
                with p.open('rb') as f:
                    while b:=f.read(8*1024*1024):out.write(b)
    manifest['merged']={'path':str(target),'bytes':target.stat().st_size,'sha256':digest(target)}
    write_json(ROOT/f'outputs/{name}-weights.json',manifest)
    print(json.dumps({'model':name,'revision':info.sha,'bytes':target.stat().st_size}),flush=True)


def public_inputs(output_root,plan_path):
    """Run locally. Neither local subject files nor any genome VCF is opened."""
    from openpyxl import load_workbook
    root=Path(output_root);root.mkdir(parents=True,exist_ok=False)
    plan=json.loads(Path(plan_path).read_text())
    # Freeze the public upstream example revision before retrieving its dataset.
    api='https://api.github.com/repos/ArcInstitute/evo2/commits/main'
    revision=json.load(urllib.request.urlopen(api,timeout=60))['sha']
    base=f'https://raw.githubusercontent.com/ArcInstitute/evo2/{revision}/notebooks/brca1/'
    sources=[]
    for filename in ['41586_2018_461_MOESM3_ESM.xlsx','GRCh37.p13_chr17.fna.gz']:
        url=base+filename
        with urllib.request.urlopen(url,timeout=120) as r:
            payload=r.read(100_000_001)
        if len(payload)>100_000_000:raise ValueError('Unexpected public dataset size')
        path=root/filename;path.write_bytes(payload)
        sources.append(dict(url=url,sha256=digest(path),bytes=len(payload)))
    with gzip.open(root/'GRCh37.p13_chr17.fna.gz','rt') as f:
        header=next(f).strip();reference=''.join(line.strip() for line in f).upper()
    if not header.startswith('>NC_000017.10'):
        raise ValueError('Expected public GRCh37 chromosome 17 reference')
    rows=load_workbook(root/'41586_2018_461_MOESM3_ESM.xlsx',read_only=True,data_only=True).active.iter_rows(values_only=True)
    next(rows);next(rows);columns=list(next(rows));data=[]
    for values in rows:
        r=dict(zip(columns,values));cls=r.get('func.class')
        if cls not in ['FUNC','LOF']:continue
        ref,alt=str(r.get('reference')).upper(),str(r.get('alt')).upper()
        if ref not in 'ACGT' or alt not in 'ACGT' or len(ref)!=1 or len(alt)!=1 or ref==alt:continue
        pos=int(r['position (hg19)']);key=f'BRCA1:{pos}:{ref}>{alt}'
        data.append(dict(id=key,pos=pos,ref=ref,alt=alt,group=cls))
    data.sort(key=lambda r:hashlib.sha256(('20260920:'+r['id']).encode()).hexdigest())
    counts={'FUNC':0,'LOF':0};seen=set();selected=[]
    for row in data:
        if counts[row['group']]>=plan['benchmark_per_class'] or row['pos'] in seen:continue
        p=row['pos']-1;sequence=reference[p-4096:p+4096]
        if len(sequence)!=8192 or set(sequence)-set('ACGT') or sequence[4096]!=row['ref']:
            raise ValueError('Benchmark assembly/reference/window mismatch')
        selected.append(dict(**row,assembly='GRCh37',window=8192,ref_sequence=sequence,
                             alt_sequence=sequence[:4096]+row['alt']+sequence[4097:]))
        seen.add(row['pos']);counts[row['group']]+=1
        if all(v==plan['benchmark_per_class'] for v in counts.values()):break
    if any(v!=plan['benchmark_per_class'] for v in counts.values()):raise ValueError('Insufficient fixed benchmark')
    for candidate in plan['candidates']:
        for size in plan['candidate_windows']:
            pos=candidate['pos'];start=pos-size//2;end=pos+size//2-1
            url=f'https://rest.ensembl.org/sequence/region/human/15:{start}..{end}:1?coord_system_version=GRCh38;content-type=text/plain'
            payload=urllib.request.urlopen(url,timeout=90).read();sequence=payload.decode().strip().upper()
            if len(sequence)!=size or set(sequence)-set('ACGT') or sequence[size//2]!=candidate['ref']:
                raise ValueError('Public GRCh38 candidate reference mismatch')
            path=root/(candidate['id']+f'-{size}-public-reference.txt');path.write_text(sequence+'\n')
            sources.append(dict(url=url,sha256=digest(path),bytes=path.stat().st_size))
            i=size//2
            selected.append(dict(**candidate,group='candidate',assembly='GRCh38',window=size,
                                 ref_sequence=sequence,alt_sequence=sequence[:i]+candidate['alt']+sequence[i+1:]))
    write_json(root/'inputs.json',dict(plan_sha256=digest(plan_path),script_sha256=digest(__file__),
        created_utc=datetime.now(timezone.utc).isoformat(),sources=sources,benchmark_counts=counts,rows=selected,
        disclosure='Public reference DNA and published mutagenesis controls only; two candidate tuples copied from the permitted report, not VCF. All alternates constructed synthetically.'))
    print(json.dumps({'benchmark_counts':counts,'candidate_windows':len(selected)-sum(counts.values())}),flush=True)


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='command',required=True)
    w=sub.add_parser('weights');w.add_argument('model')
    a=sub.add_parser('inputs');a.add_argument('output');a.add_argument('plan')
    args=p.parse_args()
    if args.command=='weights':weights(args.model)
    else:public_inputs(args.output,args.plan)
