#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.12"
# dependencies = ["markdown-it-py==3.0.0", "openpyxl==3.1.5", "defusedxml==0.7.1"]
# ///
"""Export the v18 report HTML and a filled copy of the pinned public methods template."""
import argparse
from copy import copy
from datetime import datetime,timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
from urllib.request import Request,urlopen
from zipfile import ZipFile
from markdown_it import MarkdownIt
from openpyxl import load_workbook
from openpyxl.comments import Comment
import track2_public_review_v18 as review

ROOT=Path(__file__).resolve().parents[1]
TEMPLATE_URL='https://huggingface.co/spaces/SageBio/rare-disease-real-kid-mva-hackathon-2026/resolve/'+review.REVISION+'/static/templates/methods_description_form.xlsx'
TEMPLATE_SHA='61aab080a2868a3b724e76692b83c24812112e305cd3a8b03f8f91a6b2414441'
CSS='''@page{size:A4;margin:18mm 17mm 18mm}body{font-family:Ubuntu,sans-serif;color:#142c3c;font-size:11pt;line-height:1.4;margin:0}h1{font-family:Ubuntu,sans-serif;font-weight:400;font-size:25pt;line-height:1.12;color:#123a40;margin:0 0 18pt}h2{font-family:Ubuntu,sans-serif;font-weight:400;font-size:17pt;line-height:1.2;color:#176359;margin:20pt 0 9pt;break-after:avoid}h3{font-size:13pt;break-after:avoid}p{margin:0 0 10pt;orphans:3;widows:3}a{color:#176359;text-decoration:underline;overflow-wrap:anywhere}strong{font-weight:700}table{border-collapse:collapse;width:100%;font-size:9.6pt;line-height:1.4;margin:12pt 0 16pt;table-layout:fixed}thead{display:table-header-group}tr{break-inside:avoid}th{background:#d8eee5;color:#123a40;text-align:left}td,th{border-bottom:1px solid #cad7db;padding:7pt 7pt;vertical-align:top;overflow-wrap:anywhere}th:nth-child(1){width:45%}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f1f7f5;padding:10pt;font-size:9.5pt;line-height:1.45;break-inside:avoid}code{font-family:"Nimbus Mono PS",monospace;font-size:.93em}li{margin-bottom:5pt}hr{border:0;border-top:1px solid #cad7db}*{-webkit-print-color-adjust:exact;print-color-adjust:exact}'''

def digest(data):return hashlib.sha256(data).hexdigest()
def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('output_name');a=p.parse_args()
    if not re.fullmatch('[a-z][a-z0-9-]+',a.output_name):raise ValueError('New simple output name required')
    out=ROOT/'results/feat009'/a.output_name
    if any(q.is_symlink() for q in [ROOT,ROOT/'results',out.parent,out]):raise ValueError('No symlinked paths')
    if out.exists():raise ValueError('Use a new output directory')
    report=(ROOT/'notes/track2-report-v18.md').read_text();answers=review.report_checks(report)
    template=ROOT/'results/feat009/official-requirements-independent-20260924/static/templates/methods_description_form.xlsx'
    if template.is_file():data=template.read_bytes()
    else:
        with urlopen(Request(TEMPLATE_URL,headers={'User-Agent':'Track2-public-template-export/1'}),timeout=45) as r:data=r.read(5_000_001)
    if digest(data)!=TEMPLATE_SHA:raise ValueError('Official template hash mismatch')
    with ZipFile(BytesIO(data)) as z:
        if any('vbaProject' in n or 'externalLinks/' in n for n in z.namelist()):raise ValueError('Unexpected active workbook content')
    wb=load_workbook(BytesIO(data),keep_links=False)
    if len(wb.worksheets)!=2:raise ValueError('Template sheet count changed')
    def cells(ws):
        return [(c.coordinate,c.value,c.data_type,tuple(c._style or ())) for row in ws for c in row]
    track1_before=cells(wb.worksheets[0])
    sheet=wb.worksheets[1]
    questions={f'A{n}':sheet[f'A{n}'].value for n in range(7,18)}
    if not str(questions['A9']).startswith('(required)'):raise ValueError('Wrong methods sheet')
    for cell,value in answers.items():
        sheet[cell]=value
        style=copy(sheet[cell].alignment);style.wrap_text=True;style.vertical='top';sheet[cell].alignment=style
        sheet.row_dimensions[sheet[cell].row].height=min(409,max(65,18*(len(value)//95+2)))
    sheet['A2'].comment=Comment('Historical template text is preserved. Live Track 2 instructions and organizer announcements 7/10 allow three entries; only the latest is reviewed. This workbook is a working methods artifact; submit the PDF/Markdown report.','jvv7')
    wb.active=1
    for ws in wb:
        if any(c.data_type=='f' for row in ws for c in row):raise ValueError('Unexpected formula; no formula execution permitted')
    out.mkdir()
    name='jvv7_track2_methods_v18.xlsx';wb.save(out/name)
    check=load_workbook(out/name,data_only=False,keep_links=False)
    if {k:check.worksheets[1][k].value for k in answers}!=answers:raise ValueError('Workbook answer drift')
    if {k:check.worksheets[1][k].value for k in questions}!=questions:raise ValueError('Official prompt drift')
    if cells(check.worksheets[0])!=track1_before:raise ValueError('Track 1 template sheet drift')
    md=MarkdownIt('commonmark',{'html':False}).enable('table')
    body=md.render(report)
    page='<!doctype html>\n<html lang="en"><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; base-uri \'none\'; form-action \'none\'"><title>jvv7 · Track 2 report v18</title><style>'+CSS+'</style></head><body><main>'+body+'</main></body></html>\n'
    if re.search(r'<(?:script|iframe|img|object|embed)\b',page,re.I):raise ValueError('Active or remote report element')
    (out/'jvv7_track2_report_v18.html').write_text(page)
    (out/'jvv7_track2_report_v18.md').write_text(report)
    (out/'methods-answers.json').write_text(json.dumps(answers,indent=2)+'\n')
    manifest=dict(created_utc=datetime.now(timezone.utc).isoformat(),template_url=TEMPLATE_URL,template_sha256=TEMPLATE_SHA,
        report_sha256=digest(report.encode()),script_sha256=digest(Path(__file__).read_bytes()),answer_cells=list(answers),abstract_words=len(answers['B17'].split()),
        official_questions_unchanged=True,formulas=0,formula_errors=0,track1_values_and_cell_styles_preserved=True,
        files={f.name:digest(f.read_bytes()) for f in out.iterdir()},upload_ready=False)
    (out/'documents.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps(dict(output=str(out),methods_fields=len(answers),abstract_words=manifest['abstract_words'],files=list(manifest['files']),upload_ready=False),indent=2))

if __name__=='__main__':main()
