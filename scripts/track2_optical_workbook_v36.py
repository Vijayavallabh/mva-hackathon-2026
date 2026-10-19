#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.12"
# dependencies = ["openpyxl==3.1.5", "defusedxml==0.7.1"]
# ///
"""Compare the retained public portal workbook with independently reduced BUB1B counts."""
import argparse
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',type=Path);parser.add_argument('results',type=Path)
    parser.add_argument('--output',type=Path,required=True);a=parser.parse_args()
    if a.output.exists():raise ValueError('Preserve prior audit')
    wb=load_workbook(a.workbook,read_only=True,data_only=True)
    matches=[]
    for sheet in wb:
        rows=sheet.iter_rows(values_only=True);header=next(rows)
        if 'gene_symbol' not in header:continue
        for values in rows:
            row=dict(zip(header,values))
            if row['gene_symbol']=='BUB1B':matches.append((sheet.title,row))
    if len(matches)!=1:raise ValueError('Unique BUB1B workbook row required')
    sheet,row=matches[0]
    counts=[dict(guide=row[f'sgRNA_{i}'],interphase_cells=row[f'sgRNA_{i}_interphase_cells'],
                 mitotic_cells=row[f'sgRNA_{i}_mitotic_cells']) for i in range(1,5)]
    if any(type(r[k]) is not int or r[k]<0 for r in counts for k in ['interphase_cells','mitotic_cells']):
        raise ValueError('Invalid workbook counts')
    results=json.loads(a.results.read_text())
    for arm in ['mitotic','interphase']:
        expected={r['guide']:r[arm+'_cells'] for r in counts}
        actual={r['guide']:r['cells'] for r in results['arms'][arm]['query_coverage']['BUB1B']}
        if expected!=actual:raise ValueError('Workbook/export guide counts differ')
    record=dict(passed=True,source_url='https://vesuvius.wi.mit.edu/genes.xlsx',
        workbook_sha256=hashlib.sha256(a.workbook.read_bytes()).hexdigest(),sheet=sheet,
        counts=counts,results_sha256=hashlib.sha256(a.results.read_bytes()).hexdigest(),
        script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        biological_validation=False,interpretation='Independent file/representation check; same originating experiment.')
    a.output.write_text(json.dumps(record,indent=2)+'\n');print(json.dumps(record,indent=2))


if __name__=='__main__':main()
